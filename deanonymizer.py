#!/usr/bin/env python3
"""
Anonymus — деанонимайзер: восстанавливает оригинальные значения из токенов.
Работает с .docx, .txt, .md.
"""

import json
from pathlib import Path
from tkinter import Tk, filedialog, messagebox

try:
    from docx import Document
except ImportError:
    import tkinter as tk
    root = tk.Tk()
    root.withdraw()
    messagebox.showerror("Ошибка", "Установите библиотеку: pip install python-docx")
    raise SystemExit

ROOT = Path(__file__).resolve().parent
KEYS_DIR = ROOT / "workspace" / "keys"


def deanonymize_text(text, mapping):
    replaced = 0
    for token, value in mapping.items():
        if token in text:
            text = text.replace(token, value)
            replaced += 1
    return text, replaced


def deanonymize_docx(path, mapping):
    doc = Document(path)
    replaced = 0
    for paragraph in doc.paragraphs:
        original = paragraph.text
        new_text = original
        for token, value in mapping.items():
            if token in new_text:
                new_text = new_text.replace(token, value)
        if new_text != original:
            replaced += sum(1 for token in mapping if token in original)
            if paragraph.runs:
                paragraph.runs[0].text = new_text
                for run in paragraph.runs[1:]:
                    run.text = ''
            else:
                paragraph.add_run(new_text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    original = paragraph.text
                    new_text = original
                    for token, value in mapping.items():
                        if token in new_text:
                            new_text = new_text.replace(token, value)
                    if new_text != original:
                        replaced += sum(1 for token in mapping if token in original)
                        if paragraph.runs:
                            paragraph.runs[0].text = new_text
                            for run in paragraph.runs[1:]:
                                run.text = ''
                        else:
                            paragraph.add_run(new_text)
    return doc, replaced


def main():
    root = Tk()
    root.withdraw()

    file_path = filedialog.askopenfilename(
        title="Выберите файл с заглушками",
        filetypes=[
            ("Word документы", "*.docx"),
            ("Текстовые файлы", "*.txt *.md"),
            ("Все файлы", "*.*")
        ]
    )
    if not file_path:
        return

    keys_path = filedialog.askopenfilename(
        title="Выберите файл ключей (.json)",
        initialdir=str(KEYS_DIR),
        filetypes=[("JSON файлы", "*.json")]
    )
    if not keys_path:
        return

    file_path = Path(file_path)
    keys_path = Path(keys_path)

    try:
        mapping = json.loads(keys_path.read_text(encoding='utf-8'))
    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось прочитать файл ключей:\n{e}")
        return

    total = len(mapping)
    out_path = file_path.parent / f"{file_path.stem}_deanon{file_path.suffix}"

    try:
        if file_path.suffix.lower() == '.docx':
            doc, replaced = deanonymize_docx(file_path, mapping)
            doc.save(out_path)
        else:
            text = file_path.read_text(encoding='utf-8')
            result, replaced = deanonymize_text(text, mapping)
            out_path.write_text(result, encoding='utf-8')
    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось обработать файл:\n{e}")
        return

    messagebox.showinfo(
        "Деанонимизация завершена",
        f"Восстановлено: {replaced} из {total} токенов.\n\nФайл сохранён:\n{out_path}"
    )


if __name__ == "__main__":
    main()
