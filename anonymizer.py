#!/usr/bin/env python3
"""
Anonymus — анонимайзер отчётов об археологических работах (.docx).
Заменяет чувствительные данные токенами, сохраняет ключи в JSON.
Не трогает приложения (всё после первого заголовка «Приложение»).
"""

import re
import json
from pathlib import Path
from tkinter import Tk, filedialog, messagebox

try:
    from docx import Document
except ImportError:
    import tkinter as tk
    root = tk.Tk()
    root.withdraw()
    messagebox.showerror("Ошибка", "Установите библиотеку: pip install python-docx")
    raise SystemExit

ROOT = Path(__file__).resolve().parent
ANON_DIR = ROOT / "workspace" / "anon"
KEYS_DIR = ROOT / "workspace" / "keys"

ANON_DIR.mkdir(parents=True, exist_ok=True)
KEYS_DIR.mkdir(parents=True, exist_ok=True)

OPEN_INFO_HEADING = re.compile(
    r'^(История\s+изучения|Список\s+(литературы|источников|использованных)|Библиография)',
    re.IGNORECASE
)


class Anonymizer:
    def __init__(self):
        self._mapping = {}
        self._seen = {}
        self._counters = {}

    @property
    def mapping(self):
        return self._mapping

    def _token(self, category, value):
        if value in self._seen:
            return self._seen[value]
        n = self._counters.get(category, 0) + 1
        self._counters[category] = n
        tok = f"[{category}_{n}]"
        self._mapping[tok] = value
        self._seen[value] = tok
        return tok

    def process(self, text, skip_names=False):
        # 1. Номера договоров
        def sub_contract(m):
            return f"{m.group(1)} № {self._token('ДОГОВОР', m.group(2))}"
        text = re.sub(
            r'([Дд]оговор[а-яА-Я]*)\s*№\s*([\w./-]+)',
            sub_contract, text
        )

        # 2. Названия организаций с правовой формой (аббревиатура)
        def sub_org(m):
            return f"{m.group(1)} «{self._token('ОРГ', m.group(2))}»"
        text = re.sub(
            r'(ООО|АО|ПАО|ЗАО|ФГБУН|ФГБУ|ГБУК|ФКУ|МУП|ГУП|НП)\s*«([^»]+)»',
            sub_org, text
        )

        # 2а. Организации с полной правовой формой
        text = re.sub(
            r'(обществ\w*\s+с\s+ограниченной\s+ответственностью)\s*«([^»]+)»',
            sub_org, text,
            flags=re.IGNORECASE
        )

        # 2б. Инфраструктурные объекты (НПС, ПСП, ЦПС, ЛПДС и др.)
        def sub_infra(m):
            return f"{m.group(1)} «{self._token('ОБЪЕКТ', m.group(2))}»"
        text = re.sub(
            r'(НПС|ПСП|ЦПС|ЛПДС|НКС|КНС|ГПЗ)\s*«([^»]+)»',
            sub_infra, text
        )

        # 3. Названия проектов в кавычках (с двоеточием или без, с неразрывными пробелами)
        def sub_project(m):
            return f"{m.group(1)}«{self._token('ПРОЕКТ', m.group(2))}»"
        text = re.sub(
            r'([Пп]роект[а-яА-Я]*(?:[\s\xa0]+№[\s\xa0]+[\w./-]+)?[\s\xa0]*:?[\s\xa0]*)«([^»]+)»',
            sub_project, text
        )

        # 4. Испрашиваемая территория расположена
        def sub_territory(m):
            return f"{m.group(1)}{self._token('МЕСТО', m.group(2))}"
        text = re.sub(
            r'(Испрашиваемая\s+территория\s+расположена\s+)(.+)',
            sub_territory, text
        )

        # 5. Координаты после слова «координаты» (DMS и десятичный формат)
        def sub_coords(m):
            raw = m.group(2).rstrip()
            tail = m.group(2)[len(raw):]
            if raw.endswith('.'):
                tail = '.' + tail
                raw = raw[:-1]
            return f"{m.group(1)}{self._token('КООРДИНАТЫ', raw)}{tail}"
        text = re.sub(
            r'(координат[ыа]?\s*[:\s]+)(N[0-9\s°′″\'\".,ENen]+)',
            sub_coords, text,
            flags=re.IGNORECASE
        )

        # 6. Названия месторождений и НГКМ
        def sub_site(m):
            return self._token('ОБЪЕКТ', m.group(0))
        text = re.sub(
            r'[А-ЯЁ][а-яёА-ЯЁ-]+(?:нск|ск)(?:ое|ого|ому|им|ом|ий)\s+(?:месторождени[еяю]|НГКМ)',
            sub_site, text
        )

        # 7. ФИО (пропускается в открытых разделах — истории изучения, библиографии)
        if not skip_names:
            def sub_name(m):
                return self._token('ФИО', m.group(0))
            text = re.sub(
                # 3 слова с отчеством в середине: Имя Отчество Фамилия
                r'[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:ович|евич|овна|евна)\s+[А-ЯЁ][а-яё]+'
                # 3 слова с отчеством в конце: Фамилия Имя Отчество
                r'|[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:ович|евич|овна|евна)'
                # 2 слова с отчеством: Имя Отчество (без фамилии)
                r'|[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:ович|евич|овна|евна)'
                # Инициалы + фамилия: И.О. Фамилия
                r'|[А-ЯЁ]\.[А-ЯЁ]\.\s+[А-ЯЁ][а-яё]{2,}'
                # Фамилия + инициалы: Фамилия И.О.
                r'|[А-ЯЁ][а-яё]{2,}\s+[А-ЯЁ]\.[А-ЯЁ]\.',
                sub_name, text
            )

        return text


def is_appendix(paragraph):
    text = paragraph.text.strip()
    if not re.match(r'^Приложени[еяю]\b', text, re.IGNORECASE):
        return False
    style = paragraph.style.name if paragraph.style else ''
    return 'Heading' in style or 'heading' in style or len(text) < 50


def set_paragraph_text(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ''


def anonymize_document(path):
    doc = Document(path)
    anon = Anonymizer()
    in_appendix = False
    in_open_section = False

    for paragraph in doc.paragraphs:
        if not in_appendix and is_appendix(paragraph):
            in_appendix = True
        if not in_appendix:
            style = paragraph.style.name if paragraph.style else ''
            if 'Heading' in style or 'heading' in style:
                in_open_section = bool(OPEN_INFO_HEADING.match(paragraph.text.strip()))
            original = paragraph.text
            replaced = anon.process(original, skip_names=in_open_section)
            if replaced != original:
                set_paragraph_text(paragraph, replaced)

    return doc, anon.mapping


def main():
    root = Tk()
    root.withdraw()

    src = filedialog.askopenfilename(
        title="Выберите файл отчёта (.docx)",
        filetypes=[("Word документы", "*.docx")]
    )
    if not src:
        return

    src = Path(src)

    try:
        doc, mapping = anonymize_document(src)
    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось обработать файл:\n{e}")
        return

    if not mapping:
        messagebox.showinfo("Готово", "Чувствительных данных не найдено.")
        return

    anon_path = ANON_DIR / f"{src.stem}_anon.docx"
    keys_path = KEYS_DIR / f"{src.stem}_keys.json"

    doc.save(anon_path)
    keys_path.write_text(
        json.dumps(mapping, ensure_ascii=False, indent=2),
        encoding='utf-8'
    )

    messagebox.showinfo(
        "Анонимизация завершена",
        f"Заменено значений: {len(mapping)}\n\n"
        f"Анонимизированный файл:\n{anon_path}\n\n"
        f"Ключи:\n{keys_path}"
    )


if __name__ == "__main__":
    main()
